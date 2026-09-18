"""Excel import/export helpers for the Library administrator tools.

The importer is intentionally independent from the curriculum workbook so an
administrator can maintain large reading collections without touching topics or
questions. Rows are upserted by the stable ``code`` field and the final write is
performed as one transaction.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .models import Level, LibraryItem


LIBRARY_HEADERS = [
    "codigo",
    "nivel",
    "categoria",
    "orden",
    "titulo",
    "resumen",
    "contenido",
    "fuente",
    "url_fuente",
    "publicado",
]

LIBRARY_ALIASES = {
    "code": "codigo",
    "level": "nivel",
    "level_code": "nivel",
    "category": "categoria",
    "order": "orden",
    "sort_order": "orden",
    "title": "titulo",
    "summary": "resumen",
    "content": "contenido",
    "content_text": "contenido",
    "source": "fuente",
    "source_label": "fuente",
    "source_url": "url_fuente",
    "is_published": "publicado",
}

VALID_CATEGORIES = {
    "story",
    "travel",
    "science",
    "technology",
    "society",
    "news_practice",
    "culture",
    "academic",
}

CATEGORY_LABELS = {
    "story": "Story",
    "travel": "Travel",
    "science": "Science",
    "technology": "Technology",
    "society": "Society",
    "news_practice": "News practice",
    "culture": "Culture",
    "academic": "Academic",
}


class LibraryBulkImportError(ValueError):
    pass


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _upper(value: Any) -> str:
    return _text(value).upper()


def _integer(value: Any, default: int = 0) -> int:
    if value is None or _text(value) == "":
        return default
    if isinstance(value, bool):
        raise ValueError("boolean is not an integer")
    number = float(value)
    if not number.is_integer():
        raise ValueError(f"'{value}' is not a whole number")
    return int(number)


def _boolean(value: Any, default: bool = True) -> bool:
    if value is None or _text(value) == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized in {"1", "true", "si", "sí", "yes", "y", "x"}:
        return True
    if normalized in {"0", "false", "no", "n"}:
        return False
    raise ValueError(f"'{value}' must be Sí/No or True/False")


def _normalized_header(value: Any) -> str:
    header = _text(value).lower().replace(" ", "_")
    return LIBRARY_ALIASES.get(header, header)


def parse_library_workbook(path: str | Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception as exc:
        raise LibraryBulkImportError(f"The Excel file could not be opened: {exc}") from exc

    try:
        lookup = {name.strip().upper(): name for name in workbook.sheetnames}
        actual_name = lookup.get("LECTURAS")
        if not actual_name:
            raise LibraryBulkImportError("The workbook must contain a sheet named 'LECTURAS'.")
        ws = workbook[actual_name]
        rows = ws.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise LibraryBulkImportError("Sheet 'LECTURAS' is empty.") from exc

        headers = [_normalized_header(value) for value in raw_headers]
        missing = [header for header in LIBRARY_HEADERS if header not in headers]
        if missing:
            raise LibraryBulkImportError(
                "Sheet 'LECTURAS' is missing required columns: " + ", ".join(missing)
            )

        records = []
        for excel_row, values in enumerate(rows, start=2):
            if not any(value is not None and _text(value) != "" for value in values):
                continue
            record = {
                headers[i]: values[i] if i < len(values) else None
                for i in range(len(headers))
                if headers[i]
            }
            record["_row"] = excel_row
            records.append(record)
        return records
    finally:
        workbook.close()


def _validate_row(row: dict[str, Any]):
    errors = []
    r = row["_row"]
    try:
        code = _upper(row.get("codigo"))
        level_code = _upper(row.get("nivel"))
        category = _text(row.get("categoria")).lower()
        title = _text(row.get("titulo"))
        summary = _text(row.get("resumen"))
        content = _text(row.get("contenido"))
        source_label = _text(row.get("fuente")) or None
        source_url = _text(row.get("url_fuente")) or None
        sort_order = _integer(row.get("orden"), 0)
        published = _boolean(row.get("publicado"), True)

        if not code:
            errors.append(f"LECTURAS row {r}: 'codigo' is required.")
        elif len(code) > 80:
            errors.append(f"LECTURAS row {r}: 'codigo' cannot exceed 80 characters.")
        if not level_code:
            errors.append(f"LECTURAS row {r}: 'nivel' is required.")
        if category not in VALID_CATEGORIES:
            errors.append(
                f"LECTURAS row {r}: invalid 'categoria' '{category}'. "
                f"Use one of: {', '.join(sorted(VALID_CATEGORIES))}."
            )
        if not title:
            errors.append(f"LECTURAS row {r}: 'titulo' is required.")
        elif len(title) > 200:
            errors.append(f"LECTURAS row {r}: 'titulo' cannot exceed 200 characters.")
        if not content:
            errors.append(f"LECTURAS row {r}: 'contenido' is required.")
        if source_label and len(source_label) > 160:
            errors.append(f"LECTURAS row {r}: 'fuente' cannot exceed 160 characters.")
        if source_url and len(source_url) > 500:
            errors.append(f"LECTURAS row {r}: 'url_fuente' cannot exceed 500 characters.")
        if source_url and not source_url.lower().startswith(("http://", "https://")):
            errors.append(f"LECTURAS row {r}: 'url_fuente' must start with http:// or https://.")
    except (TypeError, ValueError) as exc:
        errors.append(f"LECTURAS row {r}: {exc}.")
        return None, errors

    if errors:
        return None, errors

    return {
        "row": r,
        "code": code,
        "level_code": level_code,
        "category": category,
        "sort_order": sort_order,
        "title": title,
        "summary": summary,
        "content_text": content,
        "source_label": source_label,
        "source_url": source_url,
        "word_count": len(content.split()),
        "is_published": published,
    }, []


def validate_library_import(path: str | Path) -> dict[str, Any]:
    result = {
        "errors": [],
        "warnings": [],
        "readings": {"total": 0, "new": 0, "update": 0},
        "preview": [],
        "data": [],
    }
    try:
        rows = parse_library_workbook(path)
    except LibraryBulkImportError as exc:
        result["errors"].append(str(exc))
        return result

    data = []
    for row in rows:
        item, errors = _validate_row(row)
        result["errors"].extend(errors)
        if item:
            data.append(item)

    codes = [item["code"] for item in data]
    duplicate_codes = sorted({code for code in codes if codes.count(code) > 1})
    for code in duplicate_codes:
        result["errors"].append(f"LECTURAS: code '{code}' appears more than once in the workbook.")

    level_codes = {item["level_code"] for item in data}
    existing_level_codes = {
        row.code for row in Level.query.filter(Level.code.in_(level_codes)).all()
    } if level_codes else set()
    for item in data:
        if item["level_code"] not in existing_level_codes:
            result["errors"].append(
                f"LECTURAS row {item['row']}: level '{item['level_code']}' does not exist."
            )

    existing = {
        row.code: row
        for row in LibraryItem.query.filter(LibraryItem.code.in_(codes)).all()
    } if codes else {}

    result["readings"] = {
        "total": len(data),
        "new": sum(1 for item in data if item["code"] not in existing),
        "update": sum(1 for item in data if item["code"] in existing),
    }
    result["preview"] = [
        {
            "code": item["code"],
            "level": item["level_code"],
            "category": CATEGORY_LABELS.get(item["category"], item["category"]),
            "title": item["title"],
            "words": item["word_count"],
            "action": "Update" if item["code"] in existing else "Create",
        }
        for item in data[:25]
    ]
    result["data"] = data

    if not data:
        result["errors"].append("The workbook does not contain any reading rows to import.")

    return result


def apply_library_import(path: str | Path) -> dict[str, int]:
    validation = validate_library_import(path)
    if validation["errors"]:
        raise LibraryBulkImportError(
            "The workbook is no longer valid: " + " | ".join(validation["errors"][:5])
        )

    data = validation["data"]
    codes = {item["code"] for item in data}
    existing = {
        row.code: row
        for row in LibraryItem.query.filter(LibraryItem.code.in_(codes)).all()
    } if codes else {}

    created = updated = 0
    try:
        for item in data:
            reading = existing.get(item["code"])
            if reading is None:
                reading = LibraryItem(code=item["code"], title=item["title"], content_text=item["content_text"])
                db.session.add(reading)
                existing[item["code"]] = reading
                created += 1
            else:
                updated += 1

            reading.level_code = item["level_code"]
            reading.category = item["category"]
            reading.sort_order = item["sort_order"]
            reading.title = item["title"]
            reading.summary = item["summary"]
            reading.content_text = item["content_text"]
            reading.source_label = item["source_label"]
            reading.source_url = item["source_url"]
            reading.word_count = item["word_count"]
            reading.is_published = item["is_published"]

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return {"created_readings": created, "updated_readings": updated}


def _style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    widths = {
        "codigo": 22,
        "nivel": 10,
        "categoria": 18,
        "orden": 10,
        "titulo": 38,
        "resumen": 55,
        "contenido": 90,
        "fuente": 28,
        "url_fuente": 55,
        "publicado": 12,
    }
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LIBRARY_HEADERS))}1"
    for col, header in enumerate(LIBRARY_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths.get(header, 20)
    ws.row_dimensions[1].height = 30


def build_library_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "LECTURAS"
    help_ws = wb.create_sheet("INSTRUCCIONES")
    _style_sheet(ws)

    example = [
        "LIB_A1_099",
        "A1",
        "story",
        99,
        "A Short Morning Walk",
        "A simple A1 reading about a daily routine.",
        "Mia walks to the park every morning. She sees two dogs and says hello to her neighbor.\n\nAfter the walk, she buys bread and goes home.",
        "Original educational text",
        "",
        "Sí",
    ]
    example_fill = PatternFill("solid", fgColor="EAF2F8")
    for col, value in enumerate(example, 1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.fill = example_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[2].height = 80

    help_ws.column_dimensions["A"].width = 28
    help_ws.column_dimensions["B"].width = 105
    instructions = [
        ("CARGA MASIVA DE LECTURAS", "No cambies el nombre de la hoja LECTURAS ni los encabezados de la plantilla."),
        ("Flujo", "1) Completa la plantilla. 2) Sube el Excel en Administración > Library import. 3) Revisa la vista previa. 4) Confirma la importación."),
        ("Código", "codigo es la llave estable. Si existe, la lectura se actualiza; si no existe, se crea."),
        ("Nivel", "Debe existir en la plataforma: A1, A2, B1, B2, C1 o C2 en la configuración actual."),
        ("Categorías", "Valores permitidos: story, travel, science, technology, society, news_practice, culture, academic."),
        ("Contenido", "Puedes usar varios párrafos dentro de la misma celda. El sistema recalcula automáticamente word_count."),
        ("Fuente", "fuente y url_fuente son opcionales. Si agregas URL debe iniciar con http:// o https://."),
        ("News practice", "Usa news_practice solo para material educativo o adaptado. No lo presentes como noticia en tiempo real si no lo es."),
        ("Publicación", "publicado acepta Sí/No, True/False o 1/0. Despublicar es preferible a eliminar un texto que ya tenga vocabulario guardado por estudiantes."),
        ("Seguridad", "La validación es un dry run. La confirmación escribe todas las filas en una sola transacción; un error provoca rollback."),
    ]
    for idx, (title, body) in enumerate(instructions, 1):
        a = help_ws.cell(row=idx, column=1, value=title)
        b = help_ws.cell(row=idx, column=2, value=body)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_library_export() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "LECTURAS"
    help_ws = wb.create_sheet("INSTRUCCIONES")
    _style_sheet(ws)

    readings = LibraryItem.query.order_by(
        LibraryItem.level_code.asc(),
        LibraryItem.sort_order.asc(),
        LibraryItem.id.asc(),
    ).all()
    for row_idx, item in enumerate(readings, 2):
        values = {
            "codigo": item.code,
            "nivel": item.level_code,
            "categoria": item.category,
            "orden": item.sort_order,
            "titulo": item.title,
            "resumen": item.summary,
            "contenido": item.content_text,
            "fuente": item.source_label or "",
            "url_fuente": item.source_url or "",
            "publicado": "Sí" if item.is_published else "No",
        }
        for col_idx, header in enumerate(LIBRARY_HEADERS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values[header])
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    help_ws.column_dimensions["A"].width = 28
    help_ws.column_dimensions["B"].width = 105
    notes = [
        ("BANCO ACTUAL DE LECTURAS", "Este archivo contiene las lecturas registradas actualmente en la plataforma."),
        ("Edición masiva", "Corrige los campos que necesites y vuelve a subir este mismo archivo desde Administración > Library import."),
        ("Código", "No cambies codigo si quieres actualizar el mismo registro. Cambiarlo crea una lectura distinta."),
        ("Despublicar", "Usa publicado=No en lugar de borrar registros que ya tengan vocabulario guardado por estudiantes."),
    ]
    for idx, (title, body) in enumerate(notes, 1):
        a = help_ws.cell(row=idx, column=1, value=title)
        b = help_ws.cell(row=idx, column=2, value=body)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
