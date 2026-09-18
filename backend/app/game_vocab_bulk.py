from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .game_vocabulary import ARCADE_GAMES, KIDS_GAMES, LEVEL_ORDER, normalize_bool, normalize_code, normalize_games
from .models import GameVocabulary, GameVocabularyCategory

CATEGORY_HEADERS = ["codigo", "nombre", "visual", "orden", "kids", "estudiantes", "activo"]
WORD_HEADERS = [
    "clave", "ingles", "espanol", "visual", "categoria", "dificultad_kids",
    "nivel_cefr", "dificultad_estudiante", "kids", "estudiantes",
    "juegos_kids", "juegos_estudiantes", "orden", "activo",
]

class GameVocabularyImportError(ValueError):
    pass


def _text(v): return "" if v is None else str(v).strip()
def _int(v, default=0):
    if v is None or _text(v) == "": return default
    n=float(v)
    if not n.is_integer(): raise ValueError(f"'{v}' must be a whole number")
    return int(n)


def _sheet_rows(wb, sheet_name, headers):
    lookup={n.strip().upper():n for n in wb.sheetnames}
    actual=lookup.get(sheet_name)
    if not actual: return []
    ws=wb[actual]; rows=ws.iter_rows(values_only=True)
    try: raw=next(rows)
    except StopIteration: return []
    normalized=[_text(x).lower().replace(" ","_") for x in raw]
    missing=[x for x in headers if x not in normalized]
    if missing: raise GameVocabularyImportError(f"Sheet {sheet_name} is missing columns: {', '.join(missing)}")
    out=[]
    for n,values in enumerate(rows,start=2):
        if not any(_text(x) for x in values): continue
        row={normalized[i]: values[i] if i<len(values) else None for i in range(len(normalized)) if normalized[i]}
        row["_row"]=n; out.append(row)
    return out


def validate_game_vocabulary_import(path: str|Path):
    result={"errors":[],"warnings":[],"categories":{"total":0,"new":0,"update":0},"words":{"total":0,"new":0,"update":0},"category_data":[],"word_data":[],"preview":[]}
    try:
        wb=load_workbook(path,read_only=True,data_only=True)
        cats=_sheet_rows(wb,"CATEGORIAS_JUEGOS",CATEGORY_HEADERS)
        words=_sheet_rows(wb,"VOCABULARIO_JUEGOS",WORD_HEADERS)
        wb.close()
    except Exception as exc:
        result["errors"].append(f"The Excel file could not be validated: {exc}")
        return result
    if not cats and not words:
        result["errors"].append("Add rows to CATEGORIAS_JUEGOS or VOCABULARIO_JUEGOS.")
        return result

    cat_data=[]
    for row in cats:
        try:
            code=normalize_code(row.get("codigo"),60); name=_text(row.get("nombre")); visual=_text(row.get("visual")) or "🎮"
            if not code: raise ValueError("codigo is required")
            if not name: raise ValueError("nombre is required")
            cat_data.append({"row":row["_row"],"code":code,"name":name[:120],"visual":visual[:40],"sort_order":_int(row.get("orden"),0),"kids_enabled":normalize_bool(row.get("kids"),True),"standard_enabled":normalize_bool(row.get("estudiantes"),True),"is_active":normalize_bool(row.get("activo"),True)})
        except Exception as exc: result["errors"].append(f"CATEGORIAS_JUEGOS row {row['_row']}: {exc}.")
    if len({x['code'] for x in cat_data}) != len(cat_data): result["errors"].append("CATEGORIAS_JUEGOS contains duplicate codigo values.")

    existing_cat={r.code:r for r in GameVocabularyCategory.query.all()}
    workbook_cat={x['code'] for x in cat_data}
    word_data=[]
    for row in words:
        try:
            key=normalize_code(row.get("clave"),80); en=_text(row.get("ingles")); es=_text(row.get("espanol")); visual=_text(row.get("visual")) or "✨"; cat=normalize_code(row.get("categoria"),60)
            kd=_int(row.get("dificultad_kids"),1); sd=_int(row.get("dificultad_estudiante"),1); cefr=_text(row.get("nivel_cefr") or "A1").upper()
            if not key: raise ValueError("clave is required")
            if not en or not es: raise ValueError("ingles and espanol are required")
            if not cat or (cat not in existing_cat and cat not in workbook_cat): raise ValueError(f"categoria '{cat}' does not exist")
            if kd not in {1,2,3}: raise ValueError("dificultad_kids must be 1, 2 or 3")
            if sd not in {1,2,3,4,5}: raise ValueError("dificultad_estudiante must be 1 to 5")
            if cefr not in LEVEL_ORDER: raise ValueError("nivel_cefr must be A1, A2, B1, B2, C1 or C2")
            kg=normalize_games(row.get("juegos_kids"),KIDS_GAMES); ag=normalize_games(row.get("juegos_estudiantes"),ARCADE_GAMES)
            # Reject unknown non-empty game names instead of silently discarding them.
            for field,raw,allowed in [("juegos_kids",row.get("juegos_kids"),KIDS_GAMES),("juegos_estudiantes",row.get("juegos_estudiantes"),ARCADE_GAMES)]:
                text=_text(raw)
                if text and text.upper() not in {"ALL","TODOS","*"}:
                    raw_items={normalize_code(x,50) for x in text.replace('|',';').replace(',',';').split(';') if _text(x)}
                    unknown=sorted(raw_items-set(allowed))
                    if unknown: raise ValueError(f"{field} has unknown values: {', '.join(unknown)}")
            word_data.append({"row":row["_row"],"key":key,"english":en[:160],"spanish":es[:160],"visual":visual[:80],"category_code":cat,"kids_difficulty":kd,"cefr_level":cefr,"standard_difficulty":sd,"kids_enabled":normalize_bool(row.get("kids"),True),"standard_enabled":normalize_bool(row.get("estudiantes"),True),"kids_games":kg,"arcade_games":ag,"sort_order":_int(row.get("orden"),0),"is_active":normalize_bool(row.get("activo"),True)})
        except Exception as exc: result["errors"].append(f"VOCABULARIO_JUEGOS row {row['_row']}: {exc}.")
    if len({x['key'] for x in word_data}) != len(word_data): result["errors"].append("VOCABULARIO_JUEGOS contains duplicate clave values.")

    existing_words={r.key:r for r in GameVocabulary.query.all()}
    result["categories"]={"total":len(cat_data),"new":sum(x['code'] not in existing_cat for x in cat_data),"update":sum(x['code'] in existing_cat for x in cat_data)}
    result["words"]={"total":len(word_data),"new":sum(x['key'] not in existing_words for x in word_data),"update":sum(x['key'] in existing_words for x in word_data)}
    result["category_data"]=cat_data; result["word_data"]=word_data
    result["preview"]=[{"type":"Category","code":x['code'],"label":x['name'],"action":"Update" if x['code'] in existing_cat else "Create"} for x in cat_data[:10]] + [{"type":"Word","code":x['key'],"label":f"{x['visual']} {x['english']} → {x['spanish']}","action":"Update" if x['key'] in existing_words else "Create"} for x in word_data[:25]]
    return result


def apply_game_vocabulary_import(path: str|Path):
    v=validate_game_vocabulary_import(path)
    if v["errors"]: raise GameVocabularyImportError(" | ".join(v["errors"][:8]))
    categories={r.code:r for r in GameVocabularyCategory.query.all()}; cc=cu=wc=wu=0
    try:
        for data in v["category_data"]:
            row=categories.get(data['code'])
            if not row:
                row=GameVocabularyCategory(code=data['code'],name=data['name']); db.session.add(row); db.session.flush(); categories[data['code']]=row; cc+=1
            else: cu+=1
            for field in ["name","visual","sort_order","kids_enabled","standard_enabled","is_active"]: setattr(row,field,data[field])
        words={r.key:r for r in GameVocabulary.query.all()}
        for data in v["word_data"]:
            row=words.get(data['key'])
            if not row:
                row=GameVocabulary(key=data['key'],category_id=categories[data['category_code']].id,english=data['english'],spanish=data['spanish']); db.session.add(row); words[data['key']]=row; wc+=1
            else: wu+=1
            row.category_id=categories[data['category_code']].id
            for field in ["english","spanish","visual","kids_difficulty","cefr_level","standard_difficulty","kids_enabled","standard_enabled","kids_games","arcade_games","sort_order","is_active"]: setattr(row,field,data[field])
        db.session.commit()
    except Exception:
        db.session.rollback(); raise
    return {"created_categories":cc,"updated_categories":cu,"created_words":wc,"updated_words":wu}


def _style(ws, headers, widths):
    fill=PatternFill("solid",fgColor="4F46E5"); font=Font(color="FFFFFF",bold=True)
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}1"
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.fill=fill; c.font=font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); ws.column_dimensions[get_column_letter(i)].width=widths.get(h,18)
    ws.row_dimensions[1].height=30


def _base_workbook():
    wb=Workbook(); c=wb.active; c.title="CATEGORIAS_JUEGOS"; w=wb.create_sheet("VOCABULARIO_JUEGOS"); info=wb.create_sheet("INSTRUCCIONES")
    _style(c,CATEGORY_HEADERS,{"codigo":22,"nombre":28,"visual":12,"orden":10,"kids":12,"estudiantes":14,"activo":12})
    _style(w,WORD_HEADERS,{"clave":24,"ingles":28,"espanol":28,"visual":12,"categoria":22,"dificultad_kids":18,"nivel_cefr":14,"dificultad_estudiante":22,"kids":10,"estudiantes":14,"juegos_kids":45,"juegos_estudiantes":48,"orden":10,"activo":10})
    info.append(["GAME VOCABULARY - INSTRUCTIONS"]); info.append(["Use stable codigo/clave values. Existing values are updated; new values are created."])
    info.append(["Blank game lists or ALL mean that the word may be used in every compatible game."])
    info.append(["Kids games: "+", ".join(sorted(KIDS_GAMES))]); info.append(["Student games: "+", ".join(sorted(ARCADE_GAMES))]); info.append(["CEFR: A1, A2, B1, B2, C1, C2. Kids difficulty: 1-3. Student difficulty: 1-5."])
    info.column_dimensions['A'].width=120
    return wb


def build_game_vocabulary_template():
    wb=_base_workbook(); c=wb["CATEGORIAS_JUEGOS"]; w=wb["VOCABULARIO_JUEGOS"]
    c.append(["transport","Transport","🚗",20,"SI","SI","SI"])
    w.append(["car","car","carro","🚗","transport",1,"A1",1,"SI","SI","ALL","vocabulary_blitz;listening_sprint;word_scramble;memory_match",10,"SI"])
    buf=BytesIO(); wb.save(buf); buf.seek(0); return buf


def build_game_vocabulary_export():
    wb=_base_workbook(); c=wb["CATEGORIAS_JUEGOS"]; w=wb["VOCABULARIO_JUEGOS"]
    for row in GameVocabularyCategory.query.order_by(GameVocabularyCategory.sort_order,GameVocabularyCategory.code).all():
        c.append([row.code,row.name,row.visual,row.sort_order,"SI" if row.kids_enabled else "NO","SI" if row.standard_enabled else "NO","SI" if row.is_active else "NO"])
    for row in GameVocabulary.query.order_by(GameVocabulary.sort_order,GameVocabulary.key).all():
        w.append([row.key,row.english,row.spanish,row.visual,row.category.code if row.category else "",row.kids_difficulty,row.cefr_level,row.standard_difficulty,"SI" if row.kids_enabled else "NO","SI" if row.standard_enabled else "NO",";".join(row.kids_games or []) or "ALL",";".join(row.arcade_games or []) or "ALL",row.sort_order,"SI" if row.is_active else "NO"])
    buf=BytesIO(); wb.save(buf); buf.seek(0); return buf
