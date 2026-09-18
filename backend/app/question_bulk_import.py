"""Questions-only bulk import for the administrator interface.

Unlike ``bulk_import.py`` (which expects a workbook with both TEMAS and
PREGUNTAS sheets), this module accepts a workbook containing only a
PREGUNTAS sheet. Every question must reference a topic (``tema_codigo``)
that already exists in the database; topics are never created here.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .bulk_import import (
    QUESTION_ALIASES,
    QUESTION_HEADERS,
    VALID_QUESTION_TYPES,
    VALID_SKILLS,
    _read_sheet,
    _validate_question_row,
)
from .models import Question, Topic


class QuestionImportError(ValueError):
    """Raised when an uploaded questions-only workbook cannot be processed safely."""


def parse_question_workbook(path: str | Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception as exc:
        raise QuestionImportError(f"The Excel file could not be opened: {exc}") from exc

    try:
        return _read_sheet(workbook, "PREGUNTAS", QUESTION_ALIASES)
    finally:
        workbook.close()


def validate_question_only_import(path: str | Path) -> dict[str, Any]:
    """Validate a questions-only workbook without writing to the database."""
    result: dict[str, Any] = {
        "errors": [],
        "warnings": [],
        "questions": {"total": 0, "new": 0, "update": 0},
        "question_preview": [],
        "data": [],
    }

    try:
        raw_rows = parse_question_workbook(path)
    except QuestionImportError as exc:
        result["errors"].append(str(exc))
        return result

    question_data = []
    for row in raw_rows:
        item, errors = _validate_question_row(row)
        result["errors"].extend(errors)
        if item:
            question_data.append(item)

    question_codes = [item["code"] for item in question_data]
    duplicate_questions = sorted({code for code in question_codes if question_codes.count(code) > 1})
    for code in duplicate_questions:
        result["errors"].append(f"PREGUNTAS: code '{code}' appears more than once in the workbook.")

    referenced_topic_codes = {item["topic_code"] for item in question_data}
    existing_topics = {
        item.code: item
        for item in Topic.query.filter(Topic.code.in_(referenced_topic_codes)).all()
    } if referenced_topic_codes else {}
    for item in question_data:
        if item["topic_code"] not in existing_topics:
            result["errors"].append(
                f"PREGUNTAS row {item['row']}: topic '{item['topic_code']}' does not exist. "
                "Create the topic first (Content Manager or the topics/questions workbook)."
            )

    existing_questions = {
        item.code: item
        for item in Question.query.filter(Question.code.in_(question_codes)).all()
    } if question_codes else {}
    for item in question_data:
        existing = existing_questions.get(item["code"])
        if existing and existing.passage_id is not None:
            # Keep the passage relationship intact. A passage-linked question may
            # be edited only inside the same topic, because this workbook does
            # not carry passage metadata.
            existing_topic_code = existing.topic.code if existing.topic else None
            if existing_topic_code != item["topic_code"]:
                result["errors"].append(
                    f"PREGUNTAS row {item['row']}: '{item['code']}' belongs to a passage in "
                    f"topic '{existing_topic_code}' and cannot be moved to '{item['topic_code']}'."
                )
        elif item["question_type"] in {"reading_comprehension", "listening_comprehension"}:
            result["errors"].append(
                f"PREGUNTAS row {item['row']}: '{item['question_type']}' can only update an "
                "existing passage-linked question; creating new passages is not supported here."
            )

    result["questions"] = {
        "total": len(question_data),
        "new": sum(1 for item in question_data if item["code"] not in existing_questions),
        "update": sum(1 for item in question_data if item["code"] in existing_questions),
    }
    result["question_preview"] = [
        {
            "code": item["code"],
            "topic": item["topic_code"],
            "skill": item["skill"],
            "difficulty": item["difficulty"],
            "level_difficulty": item["level_difficulty"],
            "action": "Update" if item["code"] in existing_questions else "Create",
        }
        for item in question_data[:25]
    ]
    result["data"] = question_data

    if not question_data:
        result["errors"].append("The workbook does not contain any question rows to import.")

    return result


def apply_question_only_import(path: str | Path) -> dict[str, int]:
    """Validate again and atomically apply a previously previewed questions workbook."""
    validation = validate_question_only_import(path)
    if validation["errors"]:
        raise QuestionImportError(
            "The workbook is no longer valid: " + " | ".join(validation["errors"][:5])
        )

    data = validation["data"]
    created_questions = updated_questions = 0

    topic_codes = {item["topic_code"] for item in data}
    topics = {
        topic.code: topic
        for topic in Topic.query.filter(Topic.code.in_(topic_codes)).all()
    } if topic_codes else {}

    question_codes = {item["code"] for item in data}
    questions = {
        question.code: question
        for question in Question.query.filter(Question.code.in_(question_codes)).all()
    } if question_codes else {}

    try:
        for item in data:
            question = questions.get(item["code"])
            if question is None:
                question = Question(
                    code=item["code"],
                    topic_id=topics[item["topic_code"]].id,
                    prompt="",
                    option_a="",
                    option_b="",
                    option_c="",
                    option_d="",
                    correct_option="A",
                )
                db.session.add(question)
                questions[item["code"]] = question
                created_questions += 1
            else:
                updated_questions += 1

            question.topic_id = topics[item["topic_code"]].id
            question.sort_order = item["sort_order"]
            question.skill = item["skill"]
            question.question_type = item["question_type"]
            question.difficulty = item["difficulty"]
            question.level_difficulty = item["level_difficulty"]
            question.prompt = item["prompt"]
            question.option_a = item["option_a"]
            question.option_b = item["option_b"]
            question.option_c = item["option_c"]
            question.option_d = item["option_d"]
            question.correct_option = item["correct_option"]
            question.explanation = item["explanation"]
            question.audio_text = item["audio_text"]
            question.audio_accent = item["audio_accent"]
            question.audio_rate = item["audio_rate"]
            question.max_audio_plays = item["max_audio_plays"]
            question.is_active = item["is_active"]

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return {
        "created_questions": created_questions,
        "updated_questions": updated_questions,
    }


def build_question_only_template() -> BytesIO:
    """Create the questions-only Excel template in memory."""
    wb = Workbook()
    ws_questions = wb.active
    ws_questions.title = "PREGUNTAS"
    ws_help = wb.create_sheet("INSTRUCCIONES")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    example_fill = PatternFill("solid", fgColor="EAF2F8")

    widths = {
        "codigo": 20, "tema_codigo": 18, "orden": 10, "habilidad": 16,
        "tipo_pregunta": 28, "dificultad": 12, "dificultad_nivel": 18, "pregunta": 50,
        "opcion_a": 32, "opcion_b": 32, "opcion_c": 32, "opcion_d": 32,
        "respuesta_correcta": 18, "explicacion": 45, "audio_text": 45,
        "audio_accent": 15, "audio_rate": 12, "max_audio_plays": 18, "activa": 10,
    }
    ws_questions.freeze_panes = "A2"
    ws_questions.auto_filter.ref = f"A1:{get_column_letter(len(QUESTION_HEADERS))}1"
    for col, header in enumerate(QUESTION_HEADERS, 1):
        cell = ws_questions.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws_questions.column_dimensions[get_column_letter(col)].width = widths.get(header, 18)
    ws_questions.row_dimensions[1].height = 30

    question_example = [
        "A1_T01_Q99", "A1_T01", 99, "grammar", "multiple_choice", 1, 2,
        "Choose the correct sentence.", "She go to school.", "She goes to school.",
        "She going to school.", "She gone to school.", "B",
        "Third-person singular takes -s in the simple present.", "", "en-US", 1.0, 2, "Sí",
    ]
    for col, value in enumerate(question_example, 1):
        cell = ws_questions.cell(row=2, column=col, value=value)
        cell.fill = example_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    instructions = [
        ("PLANTILLA DE PREGUNTAS", "Esta plantilla solo importa preguntas. No crea ni modifica temas."),
        ("Flujo", "1) Completa la plantilla. 2) Súbela en Administración > Importar preguntas. 3) Revisa la validación. 4) Confirma la importación."),
        ("Tema obligatorio", "tema_codigo debe corresponder a un tema que ya exista en la plataforma. Si el tema no existe, la fila se rechaza."),
        ("Códigos", "codigo identifica cada pregunta. Si ya existe, se actualiza; si no existe, se crea."),
        ("Habilidad", "Valores permitidos: grammar, vocabulary, reading, listening."),
        ("Tipo de pregunta", "Valores: multiple_choice, reading_multiple_choice, listening_multiple_choice, error_identification. reading_comprehension y listening_comprehension solo se aceptan al actualizar preguntas que ya pertenecen a un passage."),
        ("Dificultad global", "La columna dificultad conserva el valor histórico/global del banco (1-5)."),
        ("Dificultad dentro del nivel", "dificultad_nivel va de 1 a 5 y es la variable usada por los exámenes adaptativos. Si se deja vacía, se guarda 3."),
        ("Listening", "Si tipo_pregunta = listening_multiple_choice, audio_text es obligatorio."),
        ("Sí/No", "activa acepta Sí/No, True/False o 1/0."),
        ("Preguntas con passages", "Se pueden actualizar preguntas ya ligadas a passages sin perder su passage_id, siempre que permanezcan en el mismo tema."),
        ("Importación segura", "La validación no escribe en la base de datos. Solo la confirmación final hace commit; ante un error se ejecuta rollback."),
    ]
    ws_help.column_dimensions["A"].width = 26
    ws_help.column_dimensions["B"].width = 100
    for row_idx, (title, text) in enumerate(instructions, 1):
        a = ws_help.cell(row=row_idx, column=1, value=title)
        b = ws_help.cell(row=row_idx, column=2, value=text)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if row_idx == 1:
            a.fill = header_fill
            b.fill = header_fill
            a.font = Font(color="FFFFFF", bold=True)
            b.font = Font(color="FFFFFF", bold=True)
        ws_help.row_dimensions[row_idx].height = 34 if row_idx > 1 else 28

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
