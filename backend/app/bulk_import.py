"""Bulk curriculum import helpers for the administrator interface.

Phase 1B supports an Excel workbook with two sheets:
- TEMAS
- PREGUNTAS

The import is intentionally conservative: levels must already exist and this
format preserves existing passage relationships while allowing their editable
question fields (including CEFR-relative difficulty) to be updated safely.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .models import Level, Question, Topic


VALID_SKILLS = {"grammar", "vocabulary", "reading", "listening"}
VALID_QUESTION_TYPES = {
    "multiple_choice",
    "reading_multiple_choice",
    "listening_multiple_choice",
    "error_identification",
    "reading_comprehension",
    "listening_comprehension",
}

TOPIC_HEADERS = [
    "codigo",
    "nivel",
    "orden",
    "titulo",
    "objetivo",
    "teoria",
    "ejemplos",
    "errores_comunes_json",
    "vocabulario_clave",
    "youtube_video_id",
    "publicado",
]

QUESTION_HEADERS = [
    "codigo",
    "tema_codigo",
    "orden",
    "habilidad",
    "tipo_pregunta",
    "dificultad",
    "dificultad_nivel",
    "pregunta",
    "opcion_a",
    "opcion_b",
    "opcion_c",
    "opcion_d",
    "respuesta_correcta",
    "explicacion",
    "audio_text",
    "audio_accent",
    "audio_rate",
    "max_audio_plays",
    "activa",
]

TOPIC_ALIASES = {
    "code": "codigo",
    "level": "nivel",
    "level_code": "nivel",
    "order": "orden",
    "title": "titulo",
    "objective": "objetivo",
    "theory": "teoria",
    "examples": "ejemplos",
    "common_mistakes": "errores_comunes_json",
    "key_vocabulary": "vocabulario_clave",
    "is_published": "publicado",
}

QUESTION_ALIASES = {
    "code": "codigo",
    "topic_code": "tema_codigo",
    "order": "orden",
    "skill": "habilidad",
    "question_type": "tipo_pregunta",
    "difficulty": "dificultad",
    "level_difficulty": "dificultad_nivel",
    "difficulty_in_level": "dificultad_nivel",
    "adaptive_difficulty": "dificultad_nivel",
    "prompt": "pregunta",
    "option_a": "opcion_a",
    "option_b": "opcion_b",
    "option_c": "opcion_c",
    "option_d": "opcion_d",
    "correct_option": "respuesta_correcta",
    "explanation": "explicacion",
    "is_active": "activa",
}


class BulkImportError(ValueError):
    """Raised when an uploaded workbook cannot be parsed safely."""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _upper(value: Any) -> str:
    return _text(value).upper()


def _integer(value: Any, default: int | None = None) -> int | None:
    if value is None or _text(value) == "":
        return default
    if isinstance(value, bool):
        raise ValueError("boolean is not an integer")
    number = float(value)
    if not number.is_integer():
        raise ValueError(f"'{value}' is not a whole number")
    return int(number)


def _float(value: Any, default: float | None = None) -> float | None:
    if value is None or _text(value) == "":
        return default
    return float(value)


def _boolean(value: Any, default: bool = True) -> bool:
    if value is None or _text(value) == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized in {"1", "true", "si", "sí", "yes", "y", "x"}:
        return True
    if normalized in {"0", "false", "no", "n"}:
        return False
    raise ValueError(f"'{value}' must be Sí/No or True/False")


def _simple_list(value: Any) -> list[str]:
    if value is None or _text(value) == "":
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]

    raw = _text(value)
    if raw.startswith("["):
        parsed = json.loads(raw)
        if not isinstance(parsed, list):
            raise ValueError("must be a JSON list")
        return [str(item).strip() for item in parsed if str(item).strip()]

    separator = "|" if "|" in raw else "\n"
    return [item.strip() for item in raw.split(separator) if item.strip()]


def _common_mistakes(value: Any) -> list[dict[str, str]]:
    if value is None or _text(value) == "":
        return []
    raw = _text(value)
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "must be valid JSON, e.g. "
            '[{"wrong":"She go","correct":"She goes"}]'
        ) from exc

    if not isinstance(parsed, list):
        raise ValueError("must be a JSON list")

    result = []
    for index, item in enumerate(parsed, 1):
        if not isinstance(item, dict):
            raise ValueError(f"item {index} must be an object")
        wrong = _text(item.get("wrong"))
        correct = _text(item.get("correct"))
        if not wrong or not correct:
            raise ValueError(f"item {index} requires 'wrong' and 'correct'")
        result.append({"wrong": wrong, "correct": correct})
    return result


def _normalized_header(value: Any, aliases: dict[str, str]) -> str:
    header = _text(value).lower().replace(" ", "_")
    return aliases.get(header, header)


def _read_sheet(workbook, wanted_name: str, aliases: dict[str, str]) -> list[dict[str, Any]]:
    sheet_lookup = {name.strip().upper(): name for name in workbook.sheetnames}
    actual_name = sheet_lookup.get(wanted_name.upper())
    if not actual_name:
        raise BulkImportError(f"The workbook must contain a sheet named '{wanted_name}'.")

    ws = workbook[actual_name]
    rows = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise BulkImportError(f"Sheet '{wanted_name}' is empty.") from exc

    headers = [_normalized_header(value, aliases) for value in raw_headers]
    if not any(headers):
        raise BulkImportError(f"Sheet '{wanted_name}' has no headers.")

    records = []
    for excel_row, values in enumerate(rows, start=2):
        if not any(value is not None and _text(value) != "" for value in values):
            continue
        record = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
        record["_row"] = excel_row
        records.append(record)
    return records


def parse_workbook(path: str | Path) -> dict[str, list[dict[str, Any]]]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception as exc:
        raise BulkImportError(f"The Excel file could not be opened: {exc}") from exc

    try:
        topic_rows = _read_sheet(workbook, "TEMAS", TOPIC_ALIASES)
        question_rows = _read_sheet(workbook, "PREGUNTAS", QUESTION_ALIASES)
    finally:
        workbook.close()

    return {"topics": topic_rows, "questions": question_rows}


def _validate_topic_row(row: dict[str, Any]) -> tuple[dict[str, Any] | None, list[str]]:
    errors = []
    r = row["_row"]
    try:
        code = _upper(row.get("codigo"))
        level_code = _upper(row.get("nivel"))
        title = _text(row.get("titulo"))
        theory = _text(row.get("teoria"))

        if not code:
            errors.append(f"TEMAS row {r}: 'codigo' is required.")
        elif len(code) > 40:
            errors.append(f"TEMAS row {r}: 'codigo' cannot exceed 40 characters.")
        if not level_code:
            errors.append(f"TEMAS row {r}: 'nivel' is required.")
        if not title:
            errors.append(f"TEMAS row {r}: 'titulo' is required.")
        elif len(title) > 160:
            errors.append(f"TEMAS row {r}: 'titulo' cannot exceed 160 characters.")
        if not theory:
            errors.append(f"TEMAS row {r}: 'teoria' is required.")

        order = _integer(row.get("orden"), 0)
        examples = _simple_list(row.get("ejemplos"))
        mistakes = _common_mistakes(row.get("errores_comunes_json"))
        vocabulary = _simple_list(row.get("vocabulario_clave"))
        published = _boolean(row.get("publicado"), True)
        youtube = _text(row.get("youtube_video_id")) or None
        if youtube and len(youtube) > 80:
            errors.append(f"TEMAS row {r}: 'youtube_video_id' cannot exceed 80 characters.")
    except (ValueError, TypeError, json.JSONDecodeError) as exc:
        errors.append(f"TEMAS row {r}: {exc}.")
        return None, errors

    if errors:
        return None, errors

    return {
        "row": r,
        "code": code,
        "level_code": level_code,
        "sort_order": order or 0,
        "title": title,
        "objective": _text(row.get("objetivo")),
        "theory": theory,
        "examples": examples,
        "common_mistakes": mistakes,
        "key_vocabulary": vocabulary,
        "youtube_video_id": youtube,
        "is_published": published,
    }, []


def _validate_question_row(row: dict[str, Any]) -> tuple[dict[str, Any] | None, list[str]]:
    errors = []
    r = row["_row"]
    try:
        code = _upper(row.get("codigo"))
        topic_code = _upper(row.get("tema_codigo"))
        skill = _text(row.get("habilidad")).lower()
        question_type = _text(row.get("tipo_pregunta")).lower() or "multiple_choice"
        prompt = _text(row.get("pregunta"))
        options = {
            "option_a": _text(row.get("opcion_a")),
            "option_b": _text(row.get("opcion_b")),
            "option_c": _text(row.get("opcion_c")),
            "option_d": _text(row.get("opcion_d")),
        }
        correct = _upper(row.get("respuesta_correcta"))

        if not code:
            errors.append(f"PREGUNTAS row {r}: 'codigo' is required.")
        elif len(code) > 60:
            errors.append(f"PREGUNTAS row {r}: 'codigo' cannot exceed 60 characters.")
        if not topic_code:
            errors.append(f"PREGUNTAS row {r}: 'tema_codigo' is required.")
        if skill not in VALID_SKILLS:
            errors.append(
                f"PREGUNTAS row {r}: 'habilidad' must be one of {sorted(VALID_SKILLS)}."
            )
        if question_type not in VALID_QUESTION_TYPES:
            errors.append(
                f"PREGUNTAS row {r}: unsupported 'tipo_pregunta' '{question_type}'."
            )
        if not prompt:
            errors.append(f"PREGUNTAS row {r}: 'pregunta' is required.")
        for key, value in options.items():
            if not value:
                errors.append(f"PREGUNTAS row {r}: '{key.replace('option_', 'opcion_')}' is required.")
            elif len(value) > 500:
                errors.append(f"PREGUNTAS row {r}: option text cannot exceed 500 characters.")
        if correct not in {"A", "B", "C", "D"}:
            errors.append(f"PREGUNTAS row {r}: 'respuesta_correcta' must be A, B, C or D.")

        difficulty = _integer(row.get("dificultad"))
        if difficulty is None or not 1 <= difficulty <= 5:
            errors.append(f"PREGUNTAS row {r}: 'dificultad' must be an integer from 1 to 5.")

        # Phase 1B: difficulty relative to the CEFR level of the topic.
        # Older Phase 1A workbooks remain compatible and default to 3.
        level_difficulty = _integer(row.get("dificultad_nivel"), 3)
        if level_difficulty is None or not 1 <= level_difficulty <= 5:
            errors.append(
                f"PREGUNTAS row {r}: 'dificultad_nivel' must be an integer from 1 to 5."
            )

        order = _integer(row.get("orden"), 0)
        audio_rate = _float(row.get("audio_rate"), 1.0)
        max_audio_plays = _integer(row.get("max_audio_plays"), 2)
        active = _boolean(row.get("activa"), True)
        audio_text = _text(row.get("audio_text")) or None
        audio_accent = _text(row.get("audio_accent")) or "en-US"

        if question_type == "listening_multiple_choice" and not audio_text:
            errors.append(
                f"PREGUNTAS row {r}: 'audio_text' is required for listening_multiple_choice."
            )
        if max_audio_plays is None or max_audio_plays < 1:
            errors.append(f"PREGUNTAS row {r}: 'max_audio_plays' must be at least 1.")
        if audio_rate is None or audio_rate <= 0:
            errors.append(f"PREGUNTAS row {r}: 'audio_rate' must be greater than 0.")
    except (ValueError, TypeError) as exc:
        errors.append(f"PREGUNTAS row {r}: {exc}.")
        return None, errors

    if errors:
        return None, errors

    return {
        "row": r,
        "code": code,
        "topic_code": topic_code,
        "sort_order": order or 0,
        "skill": skill,
        "question_type": question_type,
        "difficulty": difficulty,
        "level_difficulty": level_difficulty,
        "prompt": prompt,
        **options,
        "correct_option": correct,
        "explanation": _text(row.get("explicacion")),
        "audio_text": audio_text,
        "audio_accent": audio_accent,
        "audio_rate": audio_rate or 1.0,
        "max_audio_plays": max_audio_plays or 2,
        "is_active": active,
    }, []


def validate_bulk_import(path: str | Path) -> dict[str, Any]:
    """Validate a workbook and calculate a dry-run summary without writing data."""
    result: dict[str, Any] = {
        "errors": [],
        "warnings": [],
        "topics": {"total": 0, "new": 0, "update": 0},
        "questions": {"total": 0, "new": 0, "update": 0},
        "topic_preview": [],
        "question_preview": [],
        "data": {"topics": [], "questions": []},
    }

    try:
        raw = parse_workbook(path)
    except BulkImportError as exc:
        result["errors"].append(str(exc))
        return result

    topic_data = []
    question_data = []
    for row in raw["topics"]:
        item, errors = _validate_topic_row(row)
        result["errors"].extend(errors)
        if item:
            topic_data.append(item)

    for row in raw["questions"]:
        item, errors = _validate_question_row(row)
        result["errors"].extend(errors)
        if item:
            question_data.append(item)

    topic_codes = [item["code"] for item in topic_data]
    question_codes = [item["code"] for item in question_data]

    duplicate_topics = sorted({code for code in topic_codes if topic_codes.count(code) > 1})
    duplicate_questions = sorted({code for code in question_codes if question_codes.count(code) > 1})
    for code in duplicate_topics:
        result["errors"].append(f"TEMAS: code '{code}' appears more than once in the workbook.")
    for code in duplicate_questions:
        result["errors"].append(f"PREGUNTAS: code '{code}' appears more than once in the workbook.")

    level_codes = {item["level_code"] for item in topic_data}
    existing_levels = {
        item.code: item
        for item in Level.query.filter(Level.code.in_(level_codes)).all()
    } if level_codes else {}
    for item in topic_data:
        if item["level_code"] not in existing_levels:
            result["errors"].append(
                f"TEMAS row {item['row']}: level '{item['level_code']}' does not exist."
            )

    existing_topics = {
        item.code: item
        for item in Topic.query.filter(Topic.code.in_(topic_codes)).all()
    } if topic_codes else {}

    referenced_topic_codes = {item["topic_code"] for item in question_data}
    db_referenced_topics = {
        item.code: item
        for item in Topic.query.filter(Topic.code.in_(referenced_topic_codes)).all()
    } if referenced_topic_codes else {}
    upload_topic_codes = set(topic_codes)

    for item in question_data:
        if item["topic_code"] not in upload_topic_codes and item["topic_code"] not in db_referenced_topics:
            result["errors"].append(
                f"PREGUNTAS row {item['row']}: topic '{item['topic_code']}' does not exist "
                "and is not included in TEMAS."
            )

    existing_questions = {
        item.code: item
        for item in Question.query.filter(Question.code.in_(question_codes)).all()
    } if question_codes else {}
    for item in question_data:
        existing = existing_questions.get(item["code"])
        if existing and existing.passage_id is not None:
            # Keep the passage relationship intact. A passage-linked question may
            # be edited only inside the same topic, because this workbook does
            # not carry passage metadata.
            existing_topic_code = existing.topic.code if existing.topic else None
            if existing_topic_code != item["topic_code"]:
                result["errors"].append(
                    f"PREGUNTAS row {item['row']}: '{item['code']}' belongs to a passage in "
                    f"topic '{existing_topic_code}' and cannot be moved to '{item['topic_code']}'."
                )
        elif item["question_type"] in {"reading_comprehension", "listening_comprehension"}:
            result["errors"].append(
                f"PREGUNTAS row {item['row']}: '{item['question_type']}' can only update an "
                "existing passage-linked question in this phase; creating new passages is not yet supported."
            )

    result["topics"] = {
        "total": len(topic_data),
        "new": sum(1 for item in topic_data if item["code"] not in existing_topics),
        "update": sum(1 for item in topic_data if item["code"] in existing_topics),
    }
    result["questions"] = {
        "total": len(question_data),
        "new": sum(1 for item in question_data if item["code"] not in existing_questions),
        "update": sum(1 for item in question_data if item["code"] in existing_questions),
    }

    result["topic_preview"] = [
        {
            "code": item["code"],
            "level": item["level_code"],
            "title": item["title"],
            "action": "Update" if item["code"] in existing_topics else "Create",
        }
        for item in topic_data[:15]
    ]
    result["question_preview"] = [
        {
            "code": item["code"],
            "topic": item["topic_code"],
            "skill": item["skill"],
            "difficulty": item["difficulty"],
            "level_difficulty": item["level_difficulty"],
            "action": "Update" if item["code"] in existing_questions else "Create",
        }
        for item in question_data[:15]
    ]
    result["data"] = {"topics": topic_data, "questions": question_data}

    if not topic_data and not question_data:
        result["errors"].append("The workbook does not contain any topic or question rows to import.")

    return result


def apply_bulk_import(path: str | Path) -> dict[str, int]:
    """Validate again and atomically apply a previously previewed workbook."""
    validation = validate_bulk_import(path)
    if validation["errors"]:
        raise BulkImportError("The workbook is no longer valid: " + " | ".join(validation["errors"][:5]))

    data = validation["data"]
    created_topics = updated_topics = created_questions = updated_questions = 0

    level_codes = {item["level_code"] for item in data["topics"]}
    levels = {
        level.code: level
        for level in Level.query.filter(Level.code.in_(level_codes)).all()
    } if level_codes else {}

    topic_codes = {item["code"] for item in data["topics"]}
    topics = {
        topic.code: topic
        for topic in Topic.query.filter(Topic.code.in_(topic_codes)).all()
    } if topic_codes else {}

    try:
        for item in data["topics"]:
            topic = topics.get(item["code"])
            if topic is None:
                topic = Topic(code=item["code"], level_id=levels[item["level_code"]].id, theory="")
                db.session.add(topic)
                topics[item["code"]] = topic
                created_topics += 1
            else:
                updated_topics += 1

            topic.level_id = levels[item["level_code"]].id
            topic.sort_order = item["sort_order"]
            topic.title = item["title"]
            topic.objective = item["objective"]
            topic.theory = item["theory"]
            topic.examples = item["examples"]
            topic.common_mistakes = item["common_mistakes"]
            topic.key_vocabulary = item["key_vocabulary"]
            topic.youtube_video_id = item["youtube_video_id"]
            topic.is_published = item["is_published"]

        db.session.flush()

        all_referenced = {item["topic_code"] for item in data["questions"]}
        missing_in_map = all_referenced - set(topics)
        if missing_in_map:
            for topic in Topic.query.filter(Topic.code.in_(missing_in_map)).all():
                topics[topic.code] = topic

        question_codes = {item["code"] for item in data["questions"]}
        questions = {
            question.code: question
            for question in Question.query.filter(Question.code.in_(question_codes)).all()
        } if question_codes else {}

        for item in data["questions"]:
            question = questions.get(item["code"])
            if question is None:
                question = Question(
                    code=item["code"],
                    topic_id=topics[item["topic_code"]].id,
                    prompt="",
                    option_a="",
                    option_b="",
                    option_c="",
                    option_d="",
                    correct_option="A",
                )
                db.session.add(question)
                questions[item["code"]] = question
                created_questions += 1
            else:
                updated_questions += 1

            question.topic_id = topics[item["topic_code"]].id
            question.sort_order = item["sort_order"]
            question.skill = item["skill"]
            question.question_type = item["question_type"]
            question.difficulty = item["difficulty"]
            question.level_difficulty = item["level_difficulty"]
            question.prompt = item["prompt"]
            question.option_a = item["option_a"]
            question.option_b = item["option_b"]
            question.option_c = item["option_c"]
            question.option_d = item["option_d"]
            question.correct_option = item["correct_option"]
            question.explanation = item["explanation"]
            question.audio_text = item["audio_text"]
            question.audio_accent = item["audio_accent"]
            question.audio_rate = item["audio_rate"]
            question.max_audio_plays = item["max_audio_plays"]
            question.is_active = item["is_active"]

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return {
        "created_topics": created_topics,
        "updated_topics": updated_topics,
        "created_questions": created_questions,
        "updated_questions": updated_questions,
    }



def build_difficulty_review_workbook() -> BytesIO:
    """Export the current question bank for CEFR-relative difficulty curation.

    The workbook can be edited and uploaded again through the same bulk import
    screen. It intentionally contains an empty TEMAS sheet so reviewing
    question difficulty cannot accidentally rewrite lesson content.
    """
    wb = Workbook()
    ws_topics = wb.active
    ws_topics.title = "TEMAS"
    ws_questions = wb.create_sheet("PREGUNTAS")
    ws_help = wb.create_sheet("INSTRUCCIONES")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(TOPIC_HEADERS, 1):
        cell = ws_topics.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for col, header in enumerate(QUESTION_HEADERS, 1):
        cell = ws_questions.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws_questions.column_dimensions[get_column_letter(col)].width = 18

    ws_questions.freeze_panes = "A2"
    ws_questions.auto_filter.ref = f"A1:{get_column_letter(len(QUESTION_HEADERS))}1"

    questions = (
        Question.query
        .join(Topic, Question.topic_id == Topic.id)
        .join(Level, Topic.level_id == Level.id)
        .order_by(Level.sort_order.asc(), Topic.sort_order.asc(), Question.sort_order.asc(), Question.id.asc())
        .all()
    )

    for row_idx, q in enumerate(questions, 2):
        values = {
            "codigo": q.code or f"QUESTION_{q.id}",
            "tema_codigo": q.topic.code if q.topic else "",
            "orden": q.sort_order,
            "habilidad": q.skill,
            "tipo_pregunta": q.question_type,
            "dificultad": q.difficulty,
            "dificultad_nivel": q.level_difficulty,
            "pregunta": q.prompt,
            "opcion_a": q.option_a,
            "opcion_b": q.option_b,
            "opcion_c": q.option_c,
            "opcion_d": q.option_d,
            "respuesta_correcta": q.correct_option,
            "explicacion": q.explanation,
            "audio_text": q.audio_text or "",
            "audio_accent": q.audio_accent,
            "audio_rate": q.audio_rate,
            "max_audio_plays": q.max_audio_plays,
            "activa": "Sí" if q.is_active else "No",
        }
        for col_idx, header in enumerate(QUESTION_HEADERS, 1):
            cell = ws_questions.cell(row=row_idx, column=col_idx, value=values.get(header, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Emphasize the column the administrator is expected to curate.
    level_diff_col = QUESTION_HEADERS.index("dificultad_nivel") + 1
    ws_questions.column_dimensions[get_column_letter(level_diff_col)].width = 20
    for cell in ws_questions[get_column_letter(level_diff_col)]:
        if cell.row > 1:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")

    ws_help.column_dimensions["A"].width = 30
    ws_help.column_dimensions["B"].width = 100
    notes = [
        ("MATRIZ DE DIFICULTAD", "Este archivo contiene el banco actual. La hoja TEMAS queda vacía deliberadamente."),
        ("Qué editar", "Para esta revisión, modifica principalmente dificultad_nivel. Usa 1=muy fácil dentro de su nivel, 2=fácil, 3=media, 4=difícil y 5=muy difícil."),
        ("Importante", "No elimines columnas obligatorias. El mismo archivo se puede validar y volver a importar desde Administración > Carga masiva."),
        ("Passages", "Las preguntas de reading/listening que ya pertenecen a passages conservan su relación mientras no cambies tema_codigo."),
        ("Valor inicial", "La migración asigna 3 a preguntas existentes para no inventar una clasificación pedagógica automática."),
    ]
    for idx, (title, body) in enumerate(notes, 1):
        a = ws_help.cell(row=idx, column=1, value=title)
        b = ws_help.cell(row=idx, column=2, value=body)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_excel_template() -> BytesIO:
    """Create the administrator-facing Excel template in memory."""
    wb = Workbook()
    ws_topics = wb.active
    ws_topics.title = "TEMAS"
    ws_questions = wb.create_sheet("PREGUNTAS")
    ws_help = wb.create_sheet("INSTRUCCIONES")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    example_fill = PatternFill("solid", fgColor="EAF2F8")

    def style_sheet(ws, headers, widths):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = widths.get(header, 18)
        ws.row_dimensions[1].height = 30

    style_sheet(
        ws_topics,
        TOPIC_HEADERS,
        {
            "codigo": 16, "nivel": 10, "orden": 10, "titulo": 28,
            "objetivo": 38, "teoria": 55, "ejemplos": 45,
            "errores_comunes_json": 55, "vocabulario_clave": 38,
            "youtube_video_id": 20, "publicado": 12,
        },
    )
    style_sheet(
        ws_questions,
        QUESTION_HEADERS,
        {
            "codigo": 20, "tema_codigo": 18, "orden": 10, "habilidad": 16,
            "tipo_pregunta": 28, "dificultad": 12, "dificultad_nivel": 18, "pregunta": 50,
            "opcion_a": 32, "opcion_b": 32, "opcion_c": 32, "opcion_d": 32,
            "respuesta_correcta": 18, "explicacion": 45, "audio_text": 45,
            "audio_accent": 15, "audio_rate": 12, "max_audio_plays": 18, "activa": 10,
        },
    )

    topic_example = [
        "A1_T99", "A1", 99, "Sample topic", "Practice a sample structure.",
        "Write the lesson theory here.",
        "Example one|Example two|Example three",
        '[{"wrong":"She go to school.","correct":"She goes to school."}]',
        "word one|word two|word three", "", "Sí",
    ]
    question_example = [
        "A1_T99_Q01", "A1_T99", 1, "grammar", "multiple_choice", 1, 2,
        "Choose the correct sentence.", "She go to school.", "She goes to school.",
        "She going to school.", "She gone to school.", "B",
        "Third-person singular takes -s in the simple present.", "", "en-US", 1.0, 2, "Sí",
    ]
    for col, value in enumerate(topic_example, 1):
        cell = ws_topics.cell(row=2, column=col, value=value)
        cell.fill = example_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, value in enumerate(question_example, 1):
        cell = ws_questions.cell(row=2, column=col, value=value)
        cell.fill = example_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    instructions = [
        ("PLANTILLA DE CARGA MASIVA", "No cambies los nombres de las hojas TEMAS y PREGUNTAS."),
        ("Flujo", "1) Completa la plantilla. 2) Súbela en Administración > Carga masiva. 3) Revisa la validación. 4) Confirma la importación."),
        ("Códigos", "Los códigos identifican cada registro. Si un código ya existe, se actualiza; si no existe, se crea."),
        ("Niveles", "El nivel debe existir previamente en la plataforma (A1, A2, B1, B2, C1 o C2 en la instalación actual)."),
        ("Listas", "En ejemplos y vocabulario usa | para separar elementos. También se acepta una lista JSON."),
        ("Errores comunes", 'Debe ser JSON: [{"wrong":"texto incorrecto","correct":"texto correcto"}]'),
        ("Habilidad", "Valores permitidos: grammar, vocabulary, reading, listening."),
        ("Tipo de pregunta", "Valores: multiple_choice, reading_multiple_choice, listening_multiple_choice, error_identification. reading_comprehension y listening_comprehension se aceptan al actualizar preguntas que ya pertenecen a un passage."),
        ("Dificultad global", "La columna dificultad conserva el valor histórico/global del banco (1–5) para compatibilidad."),
        ("Dificultad dentro del nivel", "La nueva columna dificultad_nivel va de 1 a 5 y compara la pregunta solamente con otras de su mismo nivel CEFR. 1=más fácil dentro del nivel, 3=media, 5=más exigente. Si se deja vacía, se guarda 3."),
        ("Listening", "Si tipo_pregunta = listening_multiple_choice, audio_text es obligatorio."),
        ("Sí/No", "publicado y activa aceptan Sí/No, True/False o 1/0."),
        ("Preguntas con passages", "Se pueden actualizar preguntas ya ligadas a passages sin perder su passage_id, siempre que permanezcan en el mismo tema. La creación de passages nuevos sigue fuera del alcance de esta fase."),
        ("Importación segura", "La validación no escribe en la base de datos. Solo la confirmación final hace commit; ante un error se ejecuta rollback."),
    ]
    ws_help.column_dimensions["A"].width = 26
    ws_help.column_dimensions["B"].width = 100
    for row_idx, (title, text) in enumerate(instructions, 1):
        a = ws_help.cell(row=row_idx, column=1, value=title)
        b = ws_help.cell(row=row_idx, column=2, value=text)
        a.font = Font(bold=True)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if row_idx == 1:
            a.fill = header_fill
            b.fill = header_fill
            a.font = Font(color="FFFFFF", bold=True)
            b.font = Font(color="FFFFFF", bold=True)
        ws_help.row_dimensions[row_idx].height = 34 if row_idx > 1 else 28

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
